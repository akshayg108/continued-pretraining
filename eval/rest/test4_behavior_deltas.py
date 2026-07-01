#!/usr/bin/env python
"""Test 4 (optional) — the new kNN / LP / FT of the 60 reran jobs, and their Δ vs pre-CP.

Each reran job writes a results-json with the POST-CP metrics (post_knn_f1, post_linear_f1,
post_sft_f1); the runs use --skip-baseline, so the PRE-CP baseline comes from results.xlsx
(the 'By Method' sheet, seed-averaged). This prints, per cell: post, pre, and Δ = post - pre,
and writes eval/outputs/rest_behavior.csv.  CPU only.

Run:  python eval/rest/test4_behavior_deltas.py --results-xlsx ../results.xlsx
"""
import argparse
import csv
import json
import os
import re

import openpyxl

METH = {"simclr": "SimCLR-CP", "lejepa": "LeJEPA-CP", "diet": "DIET-CP", "mae": "MAE-CP"}


def enc_tag(fn):
    if "dinov3" in fn:
        return "DINOv3"
    if "clip" in fn:
        return "CLIP"
    if "224.mae" in fn:
        return "MAE"
    return "?"


def norm_n(v):
    """NUM_DATA cell -> canonical n string. Handles ints and 'MAX (75750)' / 'MAX(546)' strings."""
    if isinstance(v, (int, float)):
        return str(int(v))
    m = re.search(r"(\d+)", str(v))   # "MAX (75750)" -> "75750"
    return m.group(1) if m else str(v)


def load_pre(xlsx):
    """(method-CP, backbone, dataset_lower, n) -> (pre_knn, pre_lp, pre_ft) from 'By Method'."""
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    pre = {}
    for sh in wb.sheetnames:
        if sh != "By Method":
            continue
        for r in wb[sh].iter_rows(values_only=True):
            if not r or len(r) < 13 or r[1] is None:
                continue
            method, backbone, ds, n = str(r[1]), str(r[2] or ""), str(r[3] or "").lower(), norm_n(r[4])
            if isinstance(r[8], (int, float)):
                pre.setdefault((method, backbone, ds, n), (r[8], r[10], r[12]))
    return pre


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--csv", default="eval/outputs/rerun_geometry.csv")
    ap.add_argument("--results-xlsx", default="../results.xlsx")
    ap.add_argument("--out", default="eval/outputs/rest_behavior.csv")
    args = ap.parse_args()

    pre = load_pre(args.results_xlsx) if os.path.exists(args.results_xlsx) else {}
    if not pre:
        print(f"WARN: no pre-CP baselines loaded from {args.results_xlsx} (Δ will be blank)")

    rows = [r for r in csv.reader(open(args.csv)) if r and r[0].strip() not in ("", "ckpt")]
    FIELDS = ["method", "encoder", "dataset", "n", "seed",
              "post_knn", "post_lp", "post_ft", "pre_knn", "pre_lp", "pre_ft",
              "d_knn", "d_lp", "d_ft", "results_json"]
    fout = open(args.out, "w", newline="")
    w = csv.DictWriter(fout, fieldnames=FIELDS)
    w.writeheader()

    hdr = f"{'method':9}{'enc':7}{'dataset':12}{'n':>7}{'sd':>4}  {'post_knn':>9}{'d_knn':>8}{'post_lp':>9}{'d_lp':>8}{'post_ft':>9}{'d_ft':>8}"
    print(hdr)
    n_have = 0
    for r in rows:
        ck = r[0].strip()
        fn = os.path.basename(ck)
        m = re.match(r"(.+?)_(vit_base_patch16_[A-Za-z0-9_.]+)_n(\d+)_s(\d+)\.ckpt$", fn)
        ds, _, n, seed = m.groups()
        tag = enc_tag(fn)
        method = ck.split("/")[ck.split("/").index("cp") + 1]   # SimCLR / LeJEPA / DIET
        logdir = os.path.dirname(os.path.dirname(ck)).replace("/ckpts/", "/logs/")
        jpath = os.path.join(logdir, f"{tag}_{ds}_n{n}_seed{seed}.json")

        post_knn = post_lp = post_ft = None
        if os.path.exists(jpath):
            d = json.load(open(jpath))
            post_knn, post_lp, post_ft = d.get("post_knn_f1"), d.get("post_linear_f1"), d.get("post_sft_f1")
            n_have += 1
        pk, pl, pf = pre.get((METH.get(method.lower(), method), tag, ds, n), (None, None, None))

        def dd(a, b):
            return round(a - b, 4) if isinstance(a, (int, float)) and isinstance(b, (int, float)) else ""
        d_knn, d_lp, d_ft = dd(post_knn, pk), dd(post_lp, pl), dd(post_ft, pf)
        w.writerow(dict(method=method, encoder=tag, dataset=ds, n=n, seed=seed,
                        post_knn=post_knn, post_lp=post_lp, post_ft=post_ft,
                        pre_knn=pk, pre_lp=pl, pre_ft=pf, d_knn=d_knn, d_lp=d_lp, d_ft=d_ft,
                        results_json=jpath))
        f = lambda x: f"{x:.4f}" if isinstance(x, (int, float)) else "--"
        print(f"{method:9}{tag:7}{ds:12}{n:>7}{seed:>4}  "
              f"{f(post_knn):>9}{f(d_knn):>8}{f(post_lp):>9}{f(d_lp):>8}{f(post_ft):>9}{f(d_ft):>8}")

    print(f"\n{n_have}/{len(rows)} cells have a results-json so far.  wrote -> {args.out}")


if __name__ == "__main__":
    main()
